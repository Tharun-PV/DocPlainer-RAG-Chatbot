import pytest

from docplainer.config import Settings
from docplainer.ingestion import ingest_files, parse_file

from .conftest import have_ocr


def all_text(result) -> str:
    return "\n".join(c.text for c in result.chunks)


def test_digital_pdf_keeps_headings_and_pages(sample_dir):
    r = ingest_files([sample_dir / "northwind_handbook.pdf"])
    assert not r.errors
    leave = next(c for c in r.chunks if "24 days" in c.text)
    assert leave.page == 1 and leave.section.endswith("Working Hours and Leave")
    sec = next(c for c in r.chunks if "hotline 7700" in c.text)
    assert sec.page == 2 and "Information Security" in sec.section


def test_ruled_tables_become_markdown_with_captions(sample_dir):
    r = ingest_files([sample_dir / "northwind_annual_report.pdf"])
    tables = [c for c in r.chunks if c.kind == "table"]
    assert len(tables) == 3
    revenue = next(c for c in tables if "Europe" in c.text)
    assert revenue.text.startswith("Table 1: Revenue by region")
    assert "| Europe | 112.4 | 141.9 | 26.2% |" in revenue.text
    # table cells must not also leak into the prose chunks
    prose = "\n".join(c.text for c in r.chunks if c.kind == "text")
    assert "141.9" not in prose


@pytest.mark.skipif(not have_ocr(), reason="rapidocr not installed")
def test_scanned_pdf_is_ocrd_including_its_table(sample_dir):
    r = ingest_files([sample_dir / "northwind_scanned_memo.pdf"], Settings(ocr_dpi=170))
    assert not r.errors and all(c.ocr for c in r.chunks)
    text = all_text(r)
    assert "612" in text and "1.2 million" in text
    table = next(c for c in r.chunks if c.kind == "table")
    assert "297000" in table.text and table.text.count("|") >= 12


def test_scanned_pdf_without_ocr_warns_instead_of_silently_dropping(sample_dir):
    r = ingest_files([sample_dir / "northwind_scanned_memo.pdf"], Settings(ocr_enabled=False))
    assert r.errors  # nothing extractable -> reported as an error, not an empty success
    assert not r.chunks


def test_page_cap_is_reported(sample_dir):
    r = ingest_files([sample_dir / "northwind_handbook.pdf"], Settings(max_pages=1))
    assert any("first 1 of 2 pages" in w for w in r.warnings)
    assert {c.page for c in r.chunks} == {1}


def test_one_bad_file_does_not_abort_the_batch(sample_dir, tmp_path):
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"not a pdf")
    r = ingest_files([bad, sample_dir / "northwind_handbook.pdf"])
    assert "broken.pdf" in r.errors and r.chunks


def test_rejects_unsupported_and_oversized(tmp_path):
    exe = tmp_path / "x.exe"
    exe.write_bytes(b"x")
    with pytest.raises(ValueError, match="Unsupported"):
        parse_file(exe)
    big = tmp_path / "big.txt"
    big.write_bytes(b"a" * (2 * 1024 * 1024))
    with pytest.raises(ValueError, match="limit"):
        parse_file(big, Settings(max_upload_mb=1))


def test_same_content_gets_same_doc_id(sample_dir, tmp_path):
    copy = tmp_path / "renamed.pdf"
    copy.write_bytes((sample_dir / "northwind_handbook.pdf").read_bytes())
    assert parse_file(copy).doc_id == parse_file(sample_dir / "northwind_handbook.pdf").doc_id


def test_text_markdown_and_docx_and_pptx(tmp_path):
    md = tmp_path / "n.md"
    md.write_text(
        "# Title\n\nIntro paragraph with enough words.\n\n## Data\n\n"
        "| k | v |\n|---|---|\n| a | 1 |\n| b | 2 |\n"
    )
    r = ingest_files([md])
    assert {c.kind for c in r.chunks} == {"text", "table"}

    from docx import Document

    d = Document()
    d.add_heading("Plan", 1)
    d.add_paragraph("Table 1: Budget")
    t = d.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("Item", "Cost"), ("Server", "40"), ("Licence", "15")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    d.add_paragraph("Closing remarks about the budget and next steps here.")
    docx = tmp_path / "p.docx"
    d.save(docx)
    r = ingest_files([docx])
    table = next(c for c in r.chunks if c.kind == "table")
    assert table.text.startswith("Table 1: Budget") and "| Server | 40 |" in table.text

    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Roadmap"
    s.placeholders[1].text = "Ship the new retrieval pipeline in the third quarter."
    pptx = tmp_path / "r.pptx"
    prs.save(pptx)
    r = ingest_files([pptx])
    assert r.chunks[0].page == 1 and r.chunks[0].section == "Roadmap"
