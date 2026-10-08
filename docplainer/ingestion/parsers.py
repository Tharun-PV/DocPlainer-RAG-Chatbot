"""Format dispatch plus DOCX / PPTX / plain-text parsers (PDF lives in pdf.py)."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

from docplainer.config import Settings, get_settings
from docplainer.ingestion import tables
from docplainer.ingestion.pdf import OcrEngine, Progress, parse_pdf
from docplainer.schemas import Block, ParsedDocument, ParsedPage

SUPPORTED_EXTS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


def file_doc_id(path: Path) -> str:
    """Content hash: re-uploading the same file yields the same id, so indexing dedupes."""
    digest = hashlib.sha1()
    with open(path, "rb") as fh:
        for part in iter(lambda: fh.read(1 << 20), b""):
            digest.update(part)
    return digest.hexdigest()[:12]


def parse_file(
    path: str | Path,
    settings: Settings | None = None,
    ocr_engine: OcrEngine | None = None,
    progress: Progress | None = None,
    file_name: str | None = None,
) -> ParsedDocument:
    path = Path(path)
    cfg = settings or get_settings()
    ext = path.suffix.lower()
    name = file_name or path.name
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")
    if path.stat().st_size > cfg.max_upload_mb * 1024 * 1024:
        raise ValueError(f"{name} is larger than the {cfg.max_upload_mb} MB limit.")
    doc_id = file_doc_id(path)
    if ext == ".pdf":
        return parse_pdf(str(path), doc_id, name, cfg, ocr_engine, progress)
    if ext == ".docx":
        return _parse_docx(path, doc_id, name)
    if ext == ".pptx":
        return _parse_pptx(path, doc_id, name)
    return _parse_text(path, doc_id, name)


def _parse_docx(path: Path, doc_id: str, name: str) -> ParsedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    blocks: list[Block] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            match = re.match(r"(?:Heading|Title)\s*(\d*)", style, re.I)
            if match:
                level = int(match.group(1)) if match.group(1) else 1
                blocks.append(Block("heading", text, level=level))
            else:
                blocks.append(Block("paragraph", text))
        elif child.tag.endswith("}tbl"):
            rows = Table(child, document).rows
            grid = tables.clean_rows([[c.text for c in r.cells] for r in rows])
            if tables.is_meaningful(grid):
                caption = ""
                if blocks and blocks[-1].kind == "paragraph" and len(blocks[-1].text) < 160:
                    if re.match(r"(table|exhibit)\b", blocks[-1].text, re.I):
                        caption = blocks.pop().text
                blocks.append(Block("table", tables.to_markdown(grid), caption=caption))
    return ParsedDocument(doc_id, name, [ParsedPage(number=0, blocks=blocks)])


def _parse_pptx(path: Path, doc_id: str, name: str) -> ParsedDocument:
    from pptx import Presentation

    parsed = ParsedDocument(doc_id, name)
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        page = ParsedPage(number=number)
        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame and title_shape.text.strip():
            page.blocks.append(Block("heading", title_shape.text.strip(), level=1))
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                grid = tables.clean_rows([[c.text for c in r.cells] for r in shape.table.rows])
                if tables.is_meaningful(grid):
                    page.blocks.append(Block("table", tables.to_markdown(grid)))
            elif shape.has_text_frame and shape.text_frame.text.strip():
                page.blocks.append(Block("paragraph", shape.text_frame.text.strip()))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                page.blocks.append(Block("paragraph", f"Speaker notes: {notes}"))
        if page.blocks:
            parsed.pages.append(page)
    return parsed


def _parse_text(path: Path, doc_id: str, name: str) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="ignore")
    blocks: list[Block] = []
    para: list[str] = []
    rows: list[str] = []

    def flush_para() -> None:
        if para:
            blocks.append(Block("paragraph", " ".join(para)))
            para.clear()

    def flush_rows() -> None:
        if rows:
            grid = tables.clean_rows(tables.parse_markdown_table("\n".join(rows)))
            if tables.is_meaningful(grid):
                blocks.append(Block("table", tables.to_markdown(grid)))
            rows.clear()

    for line in text.splitlines():
        stripped = line.strip()
        heading = re.match(r"(#{1,6})\s+(.*)", stripped)
        if stripped.startswith("|"):
            flush_para()
            rows.append(stripped)
            continue
        flush_rows()
        if heading:
            flush_para()
            blocks.append(Block("heading", heading.group(2).strip(), level=len(heading.group(1))))
        elif not stripped:
            flush_para()
        else:
            para.append(stripped)
    flush_para()
    flush_rows()
    return ParsedDocument(doc_id, name, [ParsedPage(number=0, blocks=blocks)])
